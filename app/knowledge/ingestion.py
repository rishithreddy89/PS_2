"""Document ingestion pipeline."""

import hashlib
import re
from pathlib import Path
from typing import List, Optional

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

from app.knowledge.models import Document, DocumentChunk, DocumentType
from app.knowledge.local_embedding import get_local_embedding_service
from app.knowledge.chroma_service import get_chroma_service
from app.utils.logging.logger import get_logger

logger = get_logger(__name__)


class IngestionService:
    """Service for ingesting documents into knowledge base."""
    
    def __init__(self):
        self.embedding_service = get_local_embedding_service()
        self.chroma_service = get_chroma_service()
        self.chunk_size = 512
        self.chunk_overlap = 50
        self.ingested_docs = set()
    
    def _detect_document_type(self, file_path: str, content: str) -> DocumentType:
        """Auto-detect document type."""
        path = Path(file_path)
        name_lower = path.name.lower()
        content_lower = content.lower()
        
        if "statute" in name_lower or "regulation" in content_lower[:200]:
            return DocumentType.STATUTE
        elif "case law" in name_lower or "precedent" in content_lower[:200]:
            return DocumentType.CASE_LAW
        elif "judgment" in name_lower or "court" in content_lower[:200]:
            return DocumentType.COURT_JUDGMENT
        elif "playbook" in name_lower:
            return DocumentType.PLAYBOOK
        elif "template" in name_lower:
            return DocumentType.TEMPLATE
        elif "sample" in name_lower and "case" in name_lower:
            return DocumentType.SAMPLE_CASE
        elif "policy" in name_lower:
            return DocumentType.POLICY
        elif path.suffix == ".md":
            return DocumentType.MARKDOWN
        elif path.suffix == ".pdf":
            return DocumentType.PDF
        else:
            return DocumentType.TEXT
    
    def _read_file(self, file_path: str) -> str:
        """Read file content."""
        path = Path(file_path)
        
        if path.suffix == ".pdf":
            return self._read_pdf(file_path)
        elif path.suffix == ".docx":
            return self._read_docx(file_path)
        else:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
    
    def _read_pdf(self, file_path: str) -> str:
        """Read PDF content."""
        if PyPDF2 is None:
            raise ImportError("PyPDF2 is required for PDF support. Install with: pip install PyPDF2")
        
        try:
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages:
                    text += page.extract_text()
                logger.info("PDF parsed successfully", file=file_path, length=len(text))
                return text
        except Exception as e:
            logger.error(f"Failed to read PDF file: {e}", file=file_path)
            raise
    
    def _read_docx(self, file_path: str) -> str:
        """Read DOCX content."""
        if DocxDocument is None:
            raise ImportError("python-docx is required for DOCX support. Install with: pip install python-docx")
        
        try:
            doc = DocxDocument(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
            logger.info("DOCX parsed successfully", file=file_path, length=len(text))
            return text
        except Exception as e:
            logger.error(f"Failed to read DOCX file: {e}", file=file_path)
            raise
    
    def _chunk_text(self, text: str) -> List[str]:
        """Chunk text into smaller pieces."""
        sentences = re.split(r'(?<=[.!?])\s+', text)
        chunks = []
        current_chunk = ""
        
        for sentence in sentences:
            if len(current_chunk) + len(sentence) < self.chunk_size:
                current_chunk += sentence + " "
            else:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                current_chunk = sentence + " "
        
        if current_chunk:
            chunks.append(current_chunk.strip())
        
        return chunks
    
    def _generate_doc_id(self, file_path: str) -> str:
        """Generate unique document ID."""
        return hashlib.sha256(file_path.encode()).hexdigest()[:16]
    
    def _generate_chunk_id(self, doc_id: str, chunk_index: int) -> str:
        """Generate chunk ID."""
        return f"{doc_id}_chunk_{chunk_index}"
    
    async def ingest_file(
        self,
        file_path: str,
        document_type: Optional[DocumentType] = None,
        metadata: Optional[dict] = None,
        case_id: Optional[str] = None
    ) -> Document:
        """Ingest single file."""
        doc_id = self._generate_doc_id(file_path)
        
        if doc_id in self.ingested_docs:
            logger.info("Document already ingested", document_id=doc_id)
            return None
        
        content = self._read_file(file_path)
        detected_type = document_type or self._detect_document_type(file_path, content)
        
        metadata = metadata or {}
        if case_id:
            metadata["case_id"] = case_id
            
        document = Document(
            document_id=doc_id,
            title=Path(file_path).stem,
            content=content,
            document_type=detected_type,
            metadata=metadata,
            source=file_path
        )
        
        chunks = self._chunk_text(content)
        if not chunks:
            logger.error("Document parsing returned zero chunks", document_id=doc_id, file_path=file_path)
            raise ValueError(f"Document parsing returned zero chunks for file {file_path}")
            
        document_chunks = []
        for i, chunk in enumerate(chunks):
            chunk_meta = {
                "document_type": detected_type.value,
                "title": document.title
            }
            if case_id:
                chunk_meta["case_id"] = case_id
                
            document_chunks.append(
                DocumentChunk(
                    chunk_id=self._generate_chunk_id(doc_id, i),
                    document_id=doc_id,
                    content=chunk,
                    chunk_index=i,
                    metadata=chunk_meta
                )
            )
        
        embeddings = await self.embedding_service.embed_batch(
            [chunk.content for chunk in document_chunks]
        )
        
        collection_map = {
            DocumentType.STATUTE: "statutes",
            DocumentType.CASE_LAW: "precedents",
            DocumentType.COURT_JUDGMENT: "precedents",
            DocumentType.PLAYBOOK: "playbooks",
            DocumentType.TEMPLATE: "templates",
            DocumentType.SAMPLE_CASE: "cases"
        }
        
        if case_id:
            collection_name = "case_documents"
        else:
            collection_name = collection_map.get(detected_type, "statutes")
            
        # Verify sizes and store
        from app.core.config import get_settings
        settings = get_settings()
        collection_size_before = self.chroma_service.get_collection_size(collection_name)
        
        await self.chroma_service.insert(
            collection_name=collection_name,
            chunks=document_chunks,
            embeddings=embeddings
        )
        
        collection_size_after = self.chroma_service.get_collection_size(collection_name)
        
        self.ingested_docs.add(doc_id)
        
        chunk_ids = [c.chunk_id for c in document_chunks]
        logger.info(
            "Document ingested successfully",
            document_id=doc_id,
            case_id=case_id,
            type=detected_type.value,
            collection_name=collection_name,
            chunk_count=len(chunks),
            chunk_ids=chunk_ids,
            chunk_text_preview=chunks[0][:100] if chunks else "",
            embedding_count=len(embeddings) if embeddings else 0,
            embedding_dimension=len(embeddings[0]) if embeddings and embeddings[0] else 0,
            embedding_model=self.embedding_service.model_name,
            persist_directory=settings.chroma_persist_directory,
            collection_size_before=collection_size_before,
            collection_size_after=collection_size_after,
            metadata=[c.metadata for c in document_chunks]
        )
        
        # Verify ChromaDB storage explicitly
        stored_items = self.chroma_service.get_items_by_ids(collection_name, chunk_ids)
        stored_ids = stored_items.get("ids", [])
        if not stored_ids:
            logger.error("ChromaDB verification failed: Items not found after insert", document_id=doc_id, chunk_ids=chunk_ids)
            raise ValueError(f"ChromaDB verification failed: Inserted chunks for {doc_id} were not stored.")
        else:
            logger.info("ChromaDB storage verified", collection_exists=True, collection_name=collection_name, stored_ids=stored_ids, stored_metadata=stored_items.get("metadatas"), stored_documents_preview=[d[:50] for d in stored_items.get("documents", [])])
        
        # Validation search to prevent silent retrieval failures later
        if document_chunks and embeddings:
            test_query = chunks[0] if len(chunks) > 0 else "test query"
            test_results = await self.chroma_service.search(
                collection_name=collection_name,
                query_embedding=embeddings[0],
                n_results=1,
                where={"case_id": case_id} if case_id else None
            )
            
            logger.info(
                "Retrieval verification executed",
                generated_query_preview=test_query[:50],
                query_embedding_dimension=len(embeddings[0]),
                retrieved_ids=[r.chunk_id for r in test_results] if test_results else [],
                similarity_scores=[r.score for r in test_results] if test_results else [],
                retrieved_text=[r.content[:50] for r in test_results] if test_results else [],
                metadata=[r.metadata for r in test_results] if test_results else [],
                collection_used=collection_name,
                case_filter_used=case_id
            )
            
            if not test_results:
                logger.error("Retrieval verification failed: zero chunks returned", document_id=doc_id, case_id=case_id)
                raise ValueError(f"Verification failed: retrieval returned zero chunks for document_id {doc_id}")
            logger.info("Retrieval verification successful", retrieved_chunk_ids=[r.chunk_id for r in test_results])
        
        return document
    
    async def ingest_directory(self, directory_path: str) -> List[Document]:
        """Ingest all documents in directory."""
        path = Path(directory_path)
        documents = []
        
        for file_path in path.rglob("*"):
            if file_path.is_file() and file_path.suffix in [".txt", ".pdf", ".docx"]:
                doc = await self.ingest_file(str(file_path))
                if doc:
                    documents.append(doc)
        
        logger.info(f"Ingested directory", path=directory_path, count=len(documents))
        return documents


_ingestion_service: Optional[IngestionService] = None


def get_ingestion_service() -> IngestionService:
    """Get singleton ingestion service."""
    global _ingestion_service
    if _ingestion_service is None:
        _ingestion_service = IngestionService()
    return _ingestion_service
