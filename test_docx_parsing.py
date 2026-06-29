#!/usr/bin/env python3
"""Test DOCX parsing directly."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

def test_imports():
    """Test that all required imports work."""
    print("\n=== Testing Imports ===")
    
    try:
        from docx import Document
        print("✓ python-docx imported successfully")
    except ImportError as e:
        print(f"✗ python-docx import failed: {e}")
        return False
    
    try:
        import PyPDF2
        print("✓ PyPDF2 imported successfully")
    except ImportError as e:
        print(f"✗ PyPDF2 import failed: {e}")
        return False
    
    try:
        import sentence_transformers
        print("✓ sentence-transformers imported successfully")
    except ImportError as e:
        print(f"✗ sentence-transformers import failed: {e}")
        return False
    
    try:
        import chromadb
        print("✓ chromadb imported successfully")
    except ImportError as e:
        print(f"✗ chromadb import failed: {e}")
        return False
    
    return True


def test_docx_file():
    """Test parsing an existing DOCX file."""
    print("\n=== Testing DOCX File Parsing ===")
    
    # Find the uploaded DOCX file
    upload_dir = Path("uploads")
    if not upload_dir.exists():
        print("⊘ No uploads directory found")
        return True
    
    docx_files = list(upload_dir.rglob("*.docx"))
    if not docx_files:
        print("⊘ No DOCX files found in uploads")
        return True
    
    docx_file = docx_files[0]
    print(f"Testing file: {docx_file}")
    
    try:
        from docx import Document
        doc = Document(str(docx_file))
        
        text = "\n".join([para.text for para in doc.paragraphs])
        print(f"✓ DOCX parsed successfully")
        print(f"  - Paragraphs: {len(doc.paragraphs)}")
        print(f"  - Text length: {len(text)} characters")
        print(f"  - Preview: {text[:200]}...")
        
        return True
    except Exception as e:
        print(f"✗ DOCX parsing failed: {e}")
        import traceback
        traceback.print_exc()
        return False


def main():
    print("=" * 60)
    print("DOCX PARSING VERIFICATION")
    print("=" * 60)
    
    results = []
    
    # Test imports
    results.append(("Imports", test_imports()))
    
    # Test DOCX file
    results.append(("DOCX Parsing", test_docx_file()))
    
    # Summary
    print("\n" + "=" * 60)
    print("SUMMARY")
    print("=" * 60)
    
    passed = sum(1 for _, result in results if result)
    total = len(results)
    
    for name, result in results:
        status = "✓ PASS" if result else "✗ FAIL"
        print(f"{status}: {name}")
    
    print(f"\nTotal: {passed}/{total} tests passed")
    
    if passed == total:
        print("\n✓ All tests passed! DOCX parsing is working.")
        return 0
    else:
        print(f"\n✗ {total - passed} test(s) failed.")
        return 1


if __name__ == "__main__":
    exit_code = main()
    sys.exit(exit_code)
